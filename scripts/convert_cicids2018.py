"""Stream the CIC-IDS2018 workbook into the ignored processed-data directory."""

import argparse
import csv
import hashlib
import os
import sys
import tempfile
from pathlib import Path
from typing import Sequence
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = (
    PROJECT_ROOT / "data" / "raw" / "cicids2018" / "cicids2018.csv.xlsx"
)
DEFAULT_OUTPUT = (
    PROJECT_ROOT / "data" / "processed" / "cicids2018" / "cicids2018.csv"
)
EXPECTED_SHEET = "cicids2018"
EXPECTED_FINAL_HEADER = "Label"
HASH_CHUNK_SIZE = 1024 * 1024


class ConversionError(RuntimeError):
    """Raised when the workbook cannot safely be converted."""


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(HASH_CHUNK_SIZE), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _read_header(worksheet) -> tuple[object, ...]:
    rows = worksheet.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration as exc:
        raise ConversionError(
            f"Worksheet {EXPECTED_SHEET!r} does not contain a header row."
        ) from exc
    if not header or header[-1] != EXPECTED_FINAL_HEADER:
        final_value = header[-1] if header else None
        raise ConversionError(
            f"Expected the final worksheet header to be {EXPECTED_FINAL_HEADER!r}; "
            f"found {final_value!r}."
        )
    return tuple(header)


def _check_prerequisites(source: Path, output: Path) -> tuple[str, tuple[object, ...]]:
    if not source.is_file():
        raise FileNotFoundError(f"Source workbook does not exist: {source}")
    if source.resolve() == output.resolve():
        raise ConversionError("Source workbook and CSV destination must be different files.")
    if output.exists():
        raise FileExistsError(f"Refusing to overwrite existing destination: {output}")

    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except (OSError, BadZipFile, InvalidFileException, ValueError) as exc:
        raise ConversionError(f"Could not open Excel workbook {source}: {exc}") from exc
    try:
        if EXPECTED_SHEET not in workbook.sheetnames:
            raise ConversionError(
                f"Expected worksheet {EXPECTED_SHEET!r}; "
                f"found {workbook.sheetnames!r}."
            )
        header = _read_header(workbook[EXPECTED_SHEET])
    finally:
        workbook.close()
    return _sha256(source), header


def convert_workbook(source: Path, output: Path, *, dry_run: bool = False) -> int:
    """Convert the workbook incrementally and return rows written, including header."""
    source = source.expanduser().resolve()
    output = output.expanduser().resolve()
    source_hash_before, header = _check_prerequisites(source, output)

    if dry_run:
        print(f"Dry run: workbook and header checks passed for {source}")
        print(f"Output would be written to: {output}")
        print(f"Header columns: {len(header)}; final column: {header[-1]}")
        print("No conversion was performed.")
        return 0

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    published = False
    rows_written = 0
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8-sig",
            newline="",
            dir=output.parent,
            prefix=f".{output.name}.",
            suffix=".tmp",
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
            writer = csv.writer(temporary_file)
            workbook = load_workbook(source, read_only=True, data_only=True)
            try:
                if EXPECTED_SHEET not in workbook.sheetnames:
                    raise ConversionError(
                        f"Expected worksheet {EXPECTED_SHEET!r}; "
                        f"found {workbook.sheetnames!r}."
                    )
                worksheet = workbook[EXPECTED_SHEET]
                rows = worksheet.iter_rows(values_only=True)
                current_header = next(rows, None)
                if current_header is None or tuple(current_header) != header:
                    raise ConversionError(
                        "Worksheet header changed between prerequisite validation "
                        "and conversion."
                    )
                writer.writerow(current_header)
                rows_written = 1
                for row in rows:
                    writer.writerow(row)
                    rows_written += 1
                temporary_file.flush()
                os.fsync(temporary_file.fileno())
            finally:
                workbook.close()

        source_hash_after = _sha256(source)
        if source_hash_after != source_hash_before:
            raise ConversionError(
                "Source workbook SHA-256 changed during conversion; "
                "the CSV will not be published."
            )

        try:
            os.link(temporary_path, output)
        except FileExistsError:
            raise FileExistsError(
                f"Refusing to overwrite destination created during conversion: {output}"
            ) from None
        published = True
        temporary_path.unlink()
        temporary_path = None
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()

    if not published:
        raise ConversionError("Conversion completed without publishing the CSV.")

    print(f"Source workbook: {source}")
    print(f"Output CSV: {output}")
    print(f"Rows written (including header): {rows_written:,}")
    print(f"Output size: {output.stat().st_size:,} bytes")
    print("Source workbook remains unchanged (SHA-256 verified).")
    return rows_written


def _argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Stream the CIC-IDS2018 Excel worksheet to a new UTF-8-BOM CSV. "
            "Existing output files are never overwritten."
        )
    )
    parser.add_argument(
        "--source",
        type=Path,
        default=DEFAULT_SOURCE,
        help="source workbook (default: data/raw/cicids2018/cicids2018.csv.xlsx)",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help="new CSV path (default: data/processed/cicids2018/cicids2018.csv)",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="check workbook, sheet, header, and destination without conversion",
    )
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = _argument_parser().parse_args(argv)
    try:
        convert_workbook(args.source, args.output, dry_run=args.dry_run)
    except (
        ConversionError,
        FileNotFoundError,
        FileExistsError,
        OSError,
        BadZipFile,
        InvalidFileException,
        ValueError,
    ) as exc:
        print(f"Conversion failed: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
